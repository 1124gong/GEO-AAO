# -*- coding: utf-8 -*-
"""导入前瞻网爬取的苏州产业园数据 → mcp-server/data/parks.json
源1: 苏州市产业园数据_203条.xlsx（名称/区县/类型/地址/电话）
源2: 苏州市产业园完整名录.xlsx（面积/企业数，按名称匹配合并）
"""
import json
import openpyxl
from html import unescape

BASE = r"C:\Users\14429\WorkBuddy"
SRC1 = BASE + r"\2026-08-04-10-48-41\苏州市产业园数据_203条.xlsx"
SRC2 = BASE + r"\2026-08-07-09-19-14\苏州市产业园完整名录.xlsx"
OUT = BASE + r"\2026-10-12-09-20-37\geo-aao-park\mcp-server\data\parks.json"

# 名称关键词 → 产业方向映射
INDUSTRY_MAP = [
    ("半导体", ["半导体", "集成电路", "芯片", "纳米"]),
    ("生物医药", ["生物医药", "医药", "生命科学", "医疗", "器械"]),
    ("汽车", ["汽车", "零部件", "车联网", "智能网联"]),
    ("人工智能", ["人工智能", "AI", "数据"]),
    ("软件信息", ["软件", "信息"]),
    ("智能制造", ["智造", "制造", "装备", "机械", "机器人", "精密", "模具"]),
    ("新材料", ["新材料", "材料"]),
    ("绿色低碳", ["新能源", "环保", "低碳", "节能", "循环"]),
    ("化工", ["化工"]),
    ("服装纺织", ["服装", "纺织"]),
    ("文创", ["文创", "文化", "创意", "传媒", "动漫", "影视"]),
    ("物流", ["物流", "保税", "港口", "冷链"]),
    ("农业", ["农业", "现代农业"]),
    ("商务服务", ["商务区", "商务中心", "广场", "楼宇"]),
    ("数字产业", ["数字", "电商", "智慧"]),
    ("科技产业", ["科技", "科创", "孵化", "创新", "高新"]),
]


def infer_industry(name, park_type):
    """多标签：命中几个算几个，保持 INDUSTRY_MAP 顺序。"""
    tags = [ind for ind, kws in INDUSTRY_MAP if any(k in name for k in kws)]
    if not tags and park_type and "工业园" in str(park_type):
        return ["综合产业"]
    return tags or ["综合"]


def derive_district(district, address):
    addr = str(address or "")
    if "工业园区" in addr:
        return "苏州工业园区"
    if "相城" in addr:
        return "相城区"
    if "高新" in addr or "虎丘" in addr:
        return "虎丘区（高新区）"
    if "吴中" in addr:
        return "吴中区"
    if "吴江" in addr:
        return "吴江区"
    if "姑苏" in addr:
        return "姑苏区"
    if "常熟" in addr:
        return "常熟市"
    if "张家港" in addr:
        return "张家港市"
    if "昆山" in addr:
        return "昆山市"
    if "太仓" in addr:
        return "太仓市"
    d = str(district or "").strip()
    if d in ("虎丘区", "高新区"):
        return "虎丘区（高新区）"
    return d or "苏州市"


def fmt_area(v):
    try:
        mu = float(v)
        if mu >= 10000:
            return f"{mu/10000:.1f}万亩"
        return f"{int(mu)}亩"
    except (TypeError, ValueError):
        return ""


def fmt_corp(v):
    try:
        n = int(float(v))
        return f"{n}家" if n < 10000 else f"{n/10000:.1f}万家"
    except (TypeError, ValueError):
        return ""


# 源2：名录增强表
enhance = {}
wb2 = openpyxl.load_workbook(SRC2, read_only=True)
ws2 = wb2.active
for r in ws2.iter_rows(min_row=2, values_only=True):
    if not r or not r[1]:
        continue
    name = str(r[1]).strip()
    enhance[name] = {"面积": fmt_area(r[3]), "企业数": fmt_corp(r[4])}
wb2.close()

# 源1：主数据
parks = []
wb1 = openpyxl.load_workbook(SRC1, read_only=True)
ws1 = wb1["苏州市产业园"]
for i, r in enumerate(ws1.iter_rows(min_row=2, values_only=True), start=1):
    if not r or not r[1]:
        continue
    name = unescape(str(r[1]).strip())
    park_type = unescape(str(r[5]).strip()) if r[5] else ""
    addr = str(r[6]).strip() if r[6] else ""
    phone = str(r[7]).strip() if r[7] else ""
    extra = enhance.get(name, {})
    parks.append({
        "id": f"P{i:03d}",
        "名称": name,
        "板块": derive_district(r[4], addr),
        "产业方向": infer_industry(name, park_type),
        "载体类型": park_type or "产业园",
        "地址": addr,
        "联系电话": phone,
        "面积": extra.get("面积", ""),
        "企业数": extra.get("企业数", ""),
        "数据来源": "前瞻产业园区库（2026-08 采集）",
    })
wb1.close()

db = {
    "meta": {
        "说明": "苏州产业园台账，前瞻产业园区库采集，含地址与联系电话",
        "更新时间": "2026-10-12",
        "总数": len(parks),
    },
    "parks": parks,
    "service_system": {
        "名称": "链园链企十大服务体系（2026年度主题）",
        "方向": [
            "园区选址对接", "政策申报辅导", "金融赋能", "人才服务", "供需链接",
            "技术创新", "市场开拓", "载体运营咨询", "行业研究", "品牌传播",
        ],
        "品牌项目": ["百园万企服务行", "产业进化论峰会"],
    },
}

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(db, f, ensure_ascii=False, indent=2)

# 统计
from collections import Counter
c = Counter(p["板块"] for p in parks)
print(f"导入 {len(parks)} 家园区")
for k, v in c.most_common():
    print(f"  {k}: {v}")
matched = sum(1 for p in parks if p["面积"])
print(f"面积/企业数增强匹配: {matched}/{len(parks)}")
